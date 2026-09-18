"""Load, extract, and chunk approved RAG source documents."""

from dataclasses import dataclass, field
import hashlib
import html
from html.parser import HTMLParser
import json
from pathlib import Path
import re

import httpx

try:
    from langchain_community.document_loaders import ConfluenceLoader
except ImportError:  # pragma: no cover - optional dependency
    ConfluenceLoader = None  # type: ignore[assignment]


SUPPORTED_SUFFIXES = {".md", ".txt", ".csv", ".json", ".html", ".pdf", ".docx", ".xlsx", ".pptx", ".png", ".jpg", ".jpeg"}


@dataclass(frozen=True)
class DocumentChunk:
    chunk_id: str
    source: str
    content: str
    page: int | None = None
    metadata: dict[str, str] = field(default_factory=dict)


@dataclass(frozen=True)
class ConfluenceDocument:
    page_id: str
    title: str
    url: str
    space: str
    version: str
    updated: str
    labels: tuple[str, ...]
    content: str
    content_hash: str


def confluence_api_base(base_url: str) -> str:
    """Normalize either an Atlassian site URL or a /wiki URL to its REST root."""
    normalized = base_url.strip().rstrip("/")
    if not normalized:
        return ""
    return normalized if normalized.casefold().endswith("/wiki") else f"{normalized}/wiki"


def load_text(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("pypdf is required to ingest PDF documents") from exc

        reader = PdfReader(str(path))
        pages: list[str] = []
        for page in reader.pages:
            extracted = page.extract_text() or ""
            if extracted.strip():
                pages.append(extracted.strip())
        if pages:
            return "\n".join(pages)
        return _ocr_pdf(path)

    if path.suffix.lower() == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is required to ingest Word documents") from exc
        document = Document(str(path))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(parts)

    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to ingest Excel documents") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in workbook.worksheets:
            rows.append(f"Sheet: {sheet.title}")
            rows.extend(" | ".join(str(value) for value in row if value is not None) for row in sheet.iter_rows(values_only=True))
        return "\n".join(rows)

    if path.suffix.lower() == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("python-pptx is required to ingest PowerPoint documents") from exc
        presentation = Presentation(str(path))
        return "\n".join(
            shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
        )

    if path.suffix.lower() in {".png", ".jpg", ".jpeg"}:
        return _ocr_image(path)

    if path.suffix.lower() == ".json":
        return json.dumps(json.loads(path.read_text(encoding="utf-8", errors="ignore")), indent=2)

    return path.read_text(encoding="utf-8", errors="ignore")


def _ocr_image(path: Path) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError("OCR requires pytesseract and Pillow") from exc
    return pytesseract.image_to_string(Image.open(path)).strip()


def _ocr_pdf(path: Path) -> str:
    try:
        import fitz
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError("Scanned PDF OCR requires pymupdf, pytesseract, and Pillow") from exc

    pages: list[str] = []
    document = fitz.open(path)
    for page in document:
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
        extracted = pytesseract.image_to_string(image).strip()
        if extracted:
            pages.append(extracted)
    return "\n".join(pages)


def chunk_document(path: Path, text: str, chunk_size: int = 900, overlap: int = 120, metadata: dict[str, str] | None = None) -> list[DocumentChunk]:
    if chunk_size <= overlap:
        raise ValueError("chunk_size must be greater than overlap")

    metadata = metadata or {}
    sections = re.split(r"(?=^#{1,6}\s+)", text.strip(), flags=re.MULTILINE)
    normalized_sections = [re.sub(r"\s+", " ", section).strip() for section in sections if section.strip()]
    chunks: list[DocumentChunk] = []
    index = 0
    for section in normalized_sections:
        start = 0
        while start < len(section):
            end = min(start + chunk_size, len(section))
            if end < len(section):
                boundary = section.rfind(" ", start, end)
                if boundary > start:
                    end = boundary
            content = section[start:end].strip()
            if content:
                digest = hashlib.sha256(f"{path}:{index}:{content}".encode()).hexdigest()[:16]
                chunks.append(DocumentChunk(digest, str(path), content, metadata=metadata))
                index += 1
            if end == len(section):
                break
            start = max(end - overlap, start + 1)
    return chunks


def discover_documents(root: Path) -> list[Path]:
    discovered: list[Path] = []
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        if any(part.startswith(".") for part in path.parts):
            continue
        discovered.append(path)
    return sorted(discovered)


class _ConfluenceHTMLParser(HTMLParser):
    _BLOCK_TAGS = {"br", "p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}
    _IGNORED_TAGS = {"script", "style", "noscript", "ac:structured-macro", "ac:parameter"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.ignored_depth = 0

    def handle_starttag(self, tag: str, _attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.casefold()
        if tag in self._IGNORED_TAGS:
            self.ignored_depth += 1
        elif not self.ignored_depth and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        tag = tag.casefold()
        if tag in self._IGNORED_TAGS and self.ignored_depth:
            self.ignored_depth -= 1
        elif not self.ignored_depth and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.ignored_depth:
            self.parts.append(data)


def html_to_text(raw_html: str) -> str:
    parser = _ConfluenceHTMLParser()
    parser.feed(raw_html)
    parser.close()
    lines = [re.sub(r"[ \t\r\f\v]+", " ", line).strip() for line in "".join(parser.parts).splitlines()]
    return "\n".join(line for line in lines if line).strip()


def fetch_confluence_document(page_url: str, *, base_url: str | None = None, email: str | None = None, api_token: str | None = None) -> ConfluenceDocument:
    page_url = page_url.strip()
    if not page_url:
        raise ValueError("Confluence page URL is required")

    resolved_base = (base_url or "").rstrip("/")
    if not resolved_base:
        raise ValueError("CONFLUENCE_BASE_URL or base_url is required")

    match = re.search(r"(?:/pages/|[?&]pageId=)(?P<page_id>\d+)(?:/|&|$)", page_url)
    if not match:
        raise ValueError("Could not parse a valid Confluence page ID from the page URL")
    page_id = match.group("page_id")

    api_base = confluence_api_base(resolved_base)
    if ConfluenceLoader is not None and email and api_token:
        try:
            loader = ConfluenceLoader(url=resolved_base, username=email, api_key=api_token)
            documents = loader.load(page_ids=[page_id])
        except Exception:
            documents = []
        if documents:
            first = documents[0]
            metadata = getattr(first, "metadata", {}) or {}
            title = str(metadata.get("title") or metadata.get("page_title") or "Confluence page")
            raw_text = getattr(first, "page_content", "") or ""
            if not raw_text:
                raw_text = str(first)
            text = html_to_text(raw_text) if "<" in raw_text and ">" in raw_text else raw_text.strip()
            if not text:
                text = title
            payload = {"id": page_id, "title": title, "body": {"storage": {"value": text}}}
        else:
            payload = None

    if payload is None:
        endpoint = f"{api_base}/rest/api/content/{page_id}"
        auth = (email, api_token) if email and api_token else None
        response = httpx.get(endpoint, params={"expand": "body.storage,metadata"}, auth=auth, timeout=20.0)
        response.raise_for_status()
        payload = response.json()

    title = str(payload.get("title") or "Confluence page")
    body = payload.get("body") or {}
    storage = body.get("storage") or {}
    html = str(storage.get("value") or "")
    text = html_to_text(html)
    if not text:
        text = title
    version = str((payload.get("version") or {}).get("number") or "")
    space = str((payload.get("space") or {}).get("key") or "")
    updated = str((payload.get("version") or {}).get("when") or "")
    labels = tuple(str(item.get("name")) for item in (payload.get("metadata") or {}).get("labels", {}).get("results", []) if item.get("name"))
    digest = hashlib.sha256(text.encode("utf-8")).hexdigest()
    return ConfluenceDocument(page_id, title, page_url, space, version, updated, labels, f"# {title}\n\n{text}", digest)


def fetch_confluence_page(page_url: str, *, base_url: str | None = None, email: str | None = None, api_token: str | None = None) -> str:
    return fetch_confluence_document(page_url, base_url=base_url, email=email, api_token=api_token).content


def ingest_confluence_page(
    page_url: str,
    *,
    knowledge_root: str | Path = Path("data/knowledge_base"),
    base_url: str | None = None,
    email: str | None = None,
    api_token: str | None = None,
) -> Path:
    document = fetch_confluence_document(page_url, base_url=base_url, email=email, api_token=api_token)
    page_content = document.content
    title = document.title
    page_id_match = re.search(r"(?:/pages/|[?&]pageId=)(?P<page_id>\d+)(?:/|&|$)", page_url)
    page_id = page_id_match.group("page_id") if page_id_match else "confluence-page"
    safe_title = re.sub(r"[^A-Za-z0-9_.-]+", "_", title).strip("_") or "confluence_page"
    root = Path(knowledge_root)
    root.mkdir(parents=True, exist_ok=True)
    target = root / f"{page_id}_{safe_title}.md"
    manifest_path = root / ".confluence_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    existing = manifest.get(page_id, {})
    if existing.get("content_hash") == document.content_hash and target.exists():
        return target
    target.write_text(page_content, encoding="utf-8")
    manifest[page_id] = {
        "page_id": document.page_id,
        "source": str(target),
        "title": document.title,
        "url": document.url,
        "space": document.space,
        "version": document.version,
        "updated": document.updated,
        "labels": list(document.labels),
        "content_hash": document.content_hash,
    }
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return target


def load_chunks(root: Path) -> list[DocumentChunk]:
    chunks: list[DocumentChunk] = []
    manifest_path = root / ".confluence_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    manifest_by_source = {str(item.get("source")): item for item in manifest.values()}
    for path in discover_documents(root):
        text = load_text(path)
        if text.strip():
            source_metadata = manifest_by_source.get(str(path), {})
            metadata = {
                key: str(source_metadata[key])
                for key in ("page_id", "title", "url", "space", "version", "updated", "content_hash")
                if source_metadata.get(key)
            }
            if source_metadata.get("labels"):
                metadata["labels"] = ",".join(str(label) for label in source_metadata["labels"])
            chunks.extend(chunk_document(path, text, metadata=metadata))
    return chunks
